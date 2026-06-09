"""Tests for heroforge.db.seed.

Verifies that the workbook-backed migration populates every targeted
game-data table, that the persisted row counts match what is extracted from
the source workbook, and that re-seeding is idempotent.
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path

import openpyxl
import pytest

from heroforge.db import seed
from heroforge.db.schema import get_connection

requires_workbook = pytest.mark.skipif(
    not seed._DEFAULT_WORKBOOK.exists(),
    reason="reference workbook not available",
)

requires_data_files = pytest.mark.skipif(
    not seed._DEFAULT_DATA_DIR.exists(),
    reason="source data directory not available",
)


def _expected_count(spec: seed._WorkbookTable, rows: list[tuple]) -> int:
    """Replicate UNIQUE-key conflict handling from workbook table inserts."""
    if spec.unique_by is None:
        return len(rows)
    key_indexes = [spec.columns.index(col) for col in spec.unique_by]
    keys = {tuple(row[i] for i in key_indexes) for row in rows}
    return len(keys)


@pytest.fixture(scope="module")
def seeded_db(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Seed a temporary database once from the real workbook for the module."""
    db_path = tmp_path_factory.mktemp("seed") / "seeded.db"
    seed.seed_all(db_path=db_path)
    return db_path


@pytest.fixture(scope="module")
def source_counts() -> dict[str, int]:
    """Per-table row counts extracted from the workbook via each extractor."""
    wb = openpyxl.load_workbook(
        str(seed._DEFAULT_WORKBOOK), read_only=True, data_only=True
    )
    try:
        counts = {
            spec.table: _expected_count(spec, spec.extractor(wb))
            for spec in seed._WORKBOOK_TABLES
        }
    finally:
        wb.close()
    return counts


_TABLE_NAMES = [spec.table for spec in seed._WORKBOOK_TABLES]


@requires_workbook
class TestWorkbookSeeding:
    @pytest.mark.parametrize("table", _TABLE_NAMES)
    def test_table_is_populated(self, seeded_db: Path, table: str) -> None:
        conn = get_connection(seeded_db)
        try:
            count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        finally:
            conn.close()
        assert count > 0, f"{table} should be populated from the workbook"

    @pytest.mark.parametrize("table", _TABLE_NAMES)
    def test_count_matches_source(
        self, seeded_db: Path, source_counts: dict[str, int], table: str
    ) -> None:
        conn = get_connection(seeded_db)
        try:
            count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        finally:
            conn.close()
        assert count == source_counts[table]

    def test_reference_workbook_regression_guards(self, seeded_db: Path) -> None:
        """Independent checks to catch extractor/layout regressions."""
        conn = get_connection(seeded_db)
        try:
            assert conn.execute("SELECT COUNT(*) FROM skills").fetchone()[0] == 57
            assert conn.execute("SELECT COUNT(*) FROM armor").fetchone()[0] == 104
            assert conn.execute("SELECT COUNT(*) FROM maneuvers").fetchone()[0] == 208

            skill = conn.execute(
                "SELECT name, key_ability, armor_check_penalty "
                "FROM skills WHERE name = ?",
                ("Balance",),
            ).fetchone()
            feat = conn.execute(
                "SELECT name, type, description FROM feats WHERE name = ?",
                ("Acrobatic",),
            ).fetchone()
            armor = conn.execute(
                "SELECT name, check_penalty, arcane_spell_failure, weight "
                "FROM armor WHERE name = ?",
                ("Bark",),
            ).fetchone()
            maneuver = conn.execute(
                "SELECT name, discipline, level, type FROM maneuvers WHERE name = ?",
                ("Burning Blade",),
            ).fetchone()
        finally:
            conn.close()

        assert skill is not None
        assert dict(skill) == {
            "name": "Balance",
            "key_ability": "DEX",
            "armor_check_penalty": 1,
        }
        assert feat is not None
        assert dict(feat) == {
            "name": "Acrobatic",
            "type": "General Feats",
            "description": "+2 bonus on Jump and Tumble checks.",
        }
        assert armor is not None
        assert dict(armor) == {
            "name": "Bark",
            "check_penalty": -2,
            "arcane_spell_failure": 15,
            "weight": 15.0,
        }
        assert maneuver is not None
        assert dict(maneuver) == {
            "name": "Burning Blade",
            "discipline": "Desert Wind",
            "level": 1,
            "type": "Boost",
        }

    def test_familiar_master_abilities_seeded_from_workbook(
        self, seeded_db: Path
    ) -> None:
        """The universal familiar master benefits round-trip from the workbook.

        These are the indented lines beneath the Special Abilities → Familiar
        entry (``Class Abilities!A162:A164``).
        """
        conn = get_connection(seeded_db)
        try:
            rows = conn.execute(
                "SELECT name, description FROM familiar_master_abilities "
                "ORDER BY sort_order, id"
            ).fetchall()
        finally:
            conn.close()
        names = [r["name"] for r in rows]
        assert names == ["Alertness", "Scry on Familiar (Sp)", "Natural Link (Su)"]
        alertness = rows[0]["description"]
        assert "+2 to Spot & Listen checks" in alertness

    def test_skill_footnotes_preserve_marked_names(self, seeded_db: Path) -> None:
        conn = get_connection(seeded_db)
        try:
            row = conn.execute("""
                SELECT skill_name, raw_name, marker
                FROM skill_footnotes
                WHERE skill_name = 'Appraise'
                """).fetchone()
        finally:
            conn.close()
        assert row is not None
        assert row["raw_name"] == "Appraise¹"
        assert row["marker"] == "¹"

    def test_skill_footnote_definitions_preserve_text(self, seeded_db: Path) -> None:
        conn = get_connection(seeded_db)
        try:
            row = conn.execute("""
                SELECT source_sheet, marker, description
                FROM skill_footnote_definitions
                WHERE source_sheet = 'Character Sheet I' AND marker = '¹'
                """).fetchone()
        finally:
            conn.close()
        assert row is not None
        assert (
            row["description"]
            == "This skill can be used even if the character has zero skill ranks."
        )


@requires_workbook
class TestIdempotency:
    def test_reseeding_does_not_change_counts(self, tmp_path: Path) -> None:
        db_path = tmp_path / "idempotent.db"

        seed.seed_all(db_path=db_path)
        conn = get_connection(db_path)
        try:
            first = {
                t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                for t in _TABLE_NAMES
            }
        finally:
            conn.close()

        # Seeding a second time must not duplicate any rows.
        seed.seed_all(db_path=db_path)
        conn = get_connection(db_path)
        try:
            second = {
                t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                for t in _TABLE_NAMES
            }
        finally:
            conn.close()

        assert first == second
        assert all(v > 0 for v in second.values())


class TestMissingWorkbook:
    def test_missing_workbook_is_skipped(self, tmp_path: Path) -> None:
        """A missing workbook logs a warning and leaves tables untouched."""
        from heroforge.db.schema import initialize_database

        db_path = tmp_path / "no_workbook.db"
        conn = initialize_database(db_path)
        try:
            seed.seed_workbook(conn, tmp_path / "does_not_exist.xlsm")
            count = conn.execute("SELECT COUNT(*) FROM feats").fetchone()[0]
        finally:
            conn.close()
        assert count == 0


class TestWeaponDamageMatrix:
    """The size-aware ``weapon_damage`` matrix replaces the in-code constant."""

    @pytest.fixture(scope="class")
    def damage_db(self, tmp_path_factory: pytest.TempPathFactory) -> Path:
        from heroforge.db.schema import initialize_database

        db_path = tmp_path_factory.mktemp("weapon_damage") / "wd.db"
        conn = initialize_database(db_path)
        try:
            seed.seed_weapon_damage(conn)
        finally:
            conn.close()
        return db_path

    @requires_workbook
    def test_matrix_is_populated(self, damage_db: Path) -> None:
        conn = get_connection(damage_db)
        try:
            count = conn.execute("SELECT COUNT(*) FROM weapon_damage").fetchone()[0]
        finally:
            conn.close()
        assert count > 0

    @requires_workbook
    def test_medium_damage_matches_canonical_weapons(self, damage_db: Path) -> None:
        # Step codes come from WeaponInfo.csv Dmg1(M): Dagger=4, Longsword=6,
        # Greataxe=8, Greatsword=10.
        conn = get_connection(damage_db)
        try:
            medium = seed._load_medium_weapon_damage(conn)
        finally:
            conn.close()
        assert medium[4] == "1d4"
        assert medium[6] == "1d8"
        assert medium[8] == "1d12"
        assert medium[10] == "2d6"

    @requires_workbook
    def test_size_adjustment_follows_srd(self, damage_db: Path) -> None:
        # The Longsword progression (step 6) scales by the SRD size rules:
        # Small 1d6, Medium 1d8, Large 2d6, Huge 3d6.
        conn = get_connection(damage_db)
        try:
            row = {
                r["size"]: r["damage"]
                for r in conn.execute(
                    "SELECT size, damage FROM weapon_damage WHERE step_code = 6"
                )
            }
        finally:
            conn.close()
        assert row["Small"] == "1d6"
        assert row["Medium"] == "1d8"
        assert row["Large"] == "2d6"
        assert row["Huge"] == "3d6"

    def test_missing_workbook_leaves_table_untouched(self, tmp_path: Path) -> None:
        from heroforge.db.schema import initialize_database

        db_path = tmp_path / "no_workbook.db"
        conn = initialize_database(db_path)
        try:
            seed.seed_weapon_damage(conn, tmp_path / "does_not_exist.xlsm")
            count = conn.execute("SELECT COUNT(*) FROM weapon_damage").fetchone()[0]
        finally:
            conn.close()
        assert count == 0

    @requires_workbook
    def test_seed_weapons_decodes_medium_from_matrix(self, tmp_path: Path) -> None:
        """End-to-end: ``seed_weapons`` reads Medium damage from the matrix.

        Known step codes resolve to canonical dice; codes absent from the matrix
        fall through to the raw source value rather than being lost.
        """
        from heroforge.db.schema import initialize_database

        data_dir = tmp_path / "data"
        data_dir.mkdir()
        with (data_dir / "WeaponInfo.csv").open(
            "w", encoding="utf-8", newline=""
        ) as csv_file:
            writer = csv.writer(csv_file)
            writer.writerow(["Select A Weapon", "Dmg1(M)"])
            writer.writerow(["Longsword", "6"])  # known → 1d8
            writer.writerow(["Mystery Blade", "99"])  # unknown → passthrough

        conn = initialize_database(tmp_path / "weapons.db")
        try:
            seed.seed_weapon_damage(conn)
            seed.seed_weapons(conn, data_dir)
            rows = {
                r["name"]: r["damage_medium"]
                for r in conn.execute("SELECT name, damage_medium FROM weapons")
            }
        finally:
            conn.close()
        assert rows["Longsword"] == "1d8"
        assert rows["Mystery Blade"] == "99"


class TestHelpers:
    def test_strip_footnotes(self) -> None:
        assert seed._strip_footnotes("Appraise\u00b9") == "Appraise"
        assert seed._strip_footnotes("Craft skills\u2026\u00b9") == "Craft skills"
        assert seed._strip_footnotes("Balance ") == "Balance"

    def test_trailing_footnote_markers(self) -> None:
        assert seed._trailing_footnote_markers("Appraise\u00b9") == "\u00b9"
        assert seed._trailing_footnote_markers("Swim\u00b9\u00b9") == "\u00b9\u00b9"
        assert seed._trailing_footnote_markers("Appraise\u00b9 ") == "\u00b9"
        assert seed._trailing_footnote_markers("Use Rope") == ""

    def test_parse_skill_footnote_legend(self) -> None:
        assert seed._parse_skill_footnote_legend(
            "1   This skill can be used even if the character has zero skill ranks.\n"
            "×  This skill is a class skills for at least one of your classes.\n"
            "*   Armor check penalty, if any, applies.    **   Double the armor "
            "check penalty."
        ) == [
            (
                "\u00b9",
                "This skill can be used even if the character has zero skill ranks.",
            ),
            (
                "\u00d7",
                "This skill is a class skills for at least one of your classes.",
            ),
            ("*", "Armor check penalty, if any, applies."),
            ("**", "Double the armor check penalty."),
        ]

    def test_lstrip_separator(self) -> None:
        assert seed._lstrip_separator(" : +2 bonus") == "+2 bonus"
        assert seed._lstrip_separator("plain") == "plain"

    def test_col_out_of_range(self) -> None:
        assert seed._col(("a", "b"), "Z") == ""
        assert seed._col(("a", None), "A") == "a"
        assert seed._col(("a", None), "B") == ""

    def test_extract_spell_progression_accepts_integral_float_values(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Spells per Day"
        ws.append((None, "Wizard"))
        ws.append((1.0, 0.0, 1.0))
        ws.append((1.0, 4.0, "2.0"))
        ws.append((2.0, "5.0", None))

        rows = seed._extract_spell_progression(wb, "Spells per Day")

        assert rows == [
            ("Wizard", 1, 0, 4),
            ("Wizard", 1, 1, 2),
            ("Wizard", 2, 0, 5),
        ]

    def test_extract_spell_progression_logs_on_non_integral_data_loss(
        self, caplog: pytest.LogCaptureFixture
    ) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Spells per Day"
        ws.append((None, "Wizard"))
        ws.append((1, 0, 1))
        ws.append((1, 4, "x"))
        ws.append((1.5, 5, 3))

        caplog.set_level(logging.WARNING, logger=seed.__name__)
        rows = seed._extract_spell_progression(wb, "Spells per Day")

        assert rows == [("Wizard", 1, 0, 4)]
        assert (
            "Dropping non-integral value while extracting sheet Spells per Day: 'x'"
            in caplog.text
        )
        assert (
            "Dropping non-integral value while extracting sheet Spells per Day: 1.5"
            in caplog.text
        )
        assert "non-integral count 'x'" in caplog.text
        assert "non-integral caster level 1.5" in caplog.text


@requires_data_files
class TestDataFileSeeding:
    """Regression tests for the four ``data/`` source files.

    These previously migrated **zero** rows because the seeders looked up
    column headers that did not exist in the real files, silently leaving the
    ``weapons``, ``creatures``, ``classes`` and ``tables`` tables empty.
    """

    @pytest.fixture(scope="class")
    def data_db(self, tmp_path_factory: pytest.TempPathFactory) -> Path:
        from heroforge.db.schema import initialize_database

        db_path = tmp_path_factory.mktemp("datafiles") / "data.db"
        conn = initialize_database(db_path)
        try:
            seed.seed_weapon_damage(conn)
            seed.seed_weapons(conn, seed._DEFAULT_DATA_DIR)
            seed.seed_creatures(conn, seed._DEFAULT_DATA_DIR)
            seed.seed_tables(conn, seed._DEFAULT_DATA_DIR)
            seed.seed_classes(conn, seed._DEFAULT_DATA_DIR)
        finally:
            conn.close()
        return db_path

    @pytest.mark.parametrize(
        "table", ["weapons", "creatures", "classes", "class_skills", "tables"]
    )
    def test_table_is_populated(self, data_db: Path, table: str) -> None:
        conn = get_connection(data_db)
        try:
            count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        finally:
            conn.close()
        assert count > 0, f"{table} should be migrated from the data/ source files"

    def test_weapon_values_are_decoded(self, data_db: Path) -> None:
        conn = get_connection(data_db)
        try:
            rows = {
                r["name"]: r
                for r in conn.execute(
                    "SELECT name, damage_medium, critical, range_increment "
                    "FROM weapons WHERE name IN "
                    "('Dagger', 'Longsword', 'Greataxe', 'Greatsword')"
                )
            }
        finally:
            conn.close()
        # Damage step codes are decoded to canonical 3.5 dice.
        assert rows["Dagger"]["damage_medium"] == "1d4"
        assert rows["Longsword"]["damage_medium"] == "1d8"
        assert rows["Greataxe"]["damage_medium"] == "1d12"
        assert rows["Greatsword"]["damage_medium"] == "2d6"
        # Threat/multiplier columns are formatted into a critical string.
        assert rows["Dagger"]["critical"] == "19-20/\u00d72"
        assert rows["Greataxe"]["critical"] == "\u00d73"

    def test_class_progressions_are_named(self, data_db: Path) -> None:
        conn = get_connection(data_db)
        try:
            rows = {
                r["name"]: r
                for r in conn.execute(
                    "SELECT name, bab_progression, fort_progression, "
                    "ref_progression, will_progression, hit_die "
                    "FROM classes WHERE name IN ('Fighter', 'Wizard', 'Rogue')"
                )
            }
        finally:
            conn.close()
        assert rows["Fighter"]["bab_progression"] == "fast"
        assert rows["Fighter"]["fort_progression"] == "good"
        assert rows["Fighter"]["will_progression"] == "poor"
        assert rows["Fighter"]["hit_die"] == 10
        assert rows["Wizard"]["bab_progression"] == "slow"
        assert rows["Wizard"]["will_progression"] == "good"
        assert rows["Rogue"]["ref_progression"] == "good"

    def test_class_skills_are_marked(self, data_db: Path) -> None:
        conn = get_connection(data_db)
        try:
            fighter_skills = {
                r[0]
                for r in conn.execute(
                    "SELECT skill_name FROM class_skills WHERE class_name = 'Fighter'"
                )
            }
        finally:
            conn.close()
        assert {"Climb", "Intimidate", "Jump", "Swim"} <= fighter_skills

    def test_creature_stat_blocks_are_migrated(self, data_db: Path) -> None:
        conn = get_connection(data_db)
        try:
            wolf = conn.execute(
                "SELECT size, type, str_score, dex_score, con_score "
                "FROM creatures WHERE name = 'Wolf'"
            ).fetchone()
        finally:
            conn.close()
        assert wolf is not None
        assert wolf["type"] == "Animal"
        assert (wolf["str_score"], wolf["dex_score"], wolf["con_score"]) == (13, 15, 15)


class TestDataFileHelpers:
    @pytest.mark.parametrize(
        "header",
        [
            ["TableName", "Key", "Value"],
            ["table_name", "key", "value"],
            ["TABLE_NAME", "KEY", "VALUE"],
            ["Table Name", "Key", "Value"],
        ],
    )
    def test_seed_tables_accepts_tidy_header_variants(
        self, tmp_path: Path, header: list[str]
    ) -> None:
        from heroforge.db.schema import initialize_database

        data_dir = tmp_path / "data"
        data_dir.mkdir()
        with (data_dir / "Tables.csv").open(
            "w", encoding="utf-8", newline=""
        ) as csv_file:
            writer = csv.writer(csv_file)
            writer.writerow(header)
            writer.writerow(["AbilityScores", "STR", "10"])
            writer.writerow(["AbilityScores", "DEX", "12"])
            writer.writerow(["Alignment", "Lawful", "1"])

        conn = initialize_database(tmp_path / "tables.db")
        try:
            seed.seed_tables(conn, data_dir)
            rows = conn.execute(
                "SELECT table_name, key, value FROM tables ORDER BY table_name, key"
            ).fetchall()
        finally:
            conn.close()

        assert [dict(r) for r in rows] == [
            {"table_name": "AbilityScores", "key": "DEX", "value": "12"},
            {"table_name": "AbilityScores", "key": "STR", "value": "10"},
            {"table_name": "Alignment", "key": "Lawful", "value": "1"},
        ]

    def test_decode_weapon_damage(self) -> None:
        # The Medium-size {step_code: die} map is loaded from the seeded
        # weapon_damage table; here a representative slice stands in for it.
        medium = {4: "1d4", 6: "1d8", 10: "2d6"}
        assert seed._decode_weapon_damage("4", medium) == "1d4"
        assert seed._decode_weapon_damage("6", medium) == "1d8"
        assert seed._decode_weapon_damage("10", medium) == "2d6"
        assert seed._decode_weapon_damage("", medium) == ""
        # Codes absent from the matrix fall through to the raw source value.
        assert seed._decode_weapon_damage("99", medium) == "99"
        # Non-numeric source values (cell references) are preserved verbatim.
        assert seed._decode_weapon_damage("ref:Unarmed", medium) == "ref:Unarmed"

    def test_format_weapon_critical(self) -> None:
        assert seed._format_weapon_critical("20", "2") == "\u00d72"
        assert seed._format_weapon_critical("19", "2") == "19-20/\u00d72"
        assert seed._format_weapon_critical("", "3") == "\u00d73"

    def test_bab_progression(self) -> None:
        assert seed._bab_progression(1) == "fast"
        assert seed._bab_progression(0.75) == "medium"
        assert seed._bab_progression(0.5) == "slow"
        assert seed._bab_progression(None) == ""

    def test_save_progression(self) -> None:
        assert seed._save_progression(0.5) == "good"
        assert seed._save_progression(0.34) == "poor"
        assert seed._save_progression("") == ""
